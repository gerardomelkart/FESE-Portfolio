"""Independent synthetic monthly operations demo; no institutional code or inputs."""
import argparse
import csv
import hashlib
import json
import random
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

CENTERS = ('Aurora', 'Boreal', 'Coral')
CODES = ('A01', 'A02', 'B01', 'B02')
FIELDS = ('period', 'center', 'code', 'total')


def validate_period(value):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', value):
        raise ValueError('Period must use YYYY-MM')
    datetime.strptime(value, '%Y-%m')
    return value


def read_input(path, period):
    validate_period(period)
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            if tuple(reader.fieldnames or ()) != FIELDS:
                raise ValueError(f'Expected columns: {FIELDS}')
            rows = list(reader)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = iter(workbook.active.values)
            if tuple(next(values, ())) != FIELDS:
                raise ValueError(f'Expected columns: {FIELDS}')
            rows = [dict(zip(FIELDS, row)) for row in values if any(v is not None for v in row)]
        finally:
            workbook.close()
    else:
        raise ValueError('Input must be CSV or XLSX')
    clean, seen = [], set()
    for row in rows:
        if set(row) != set(FIELDS) or None in row.values():
            raise ValueError('Missing or extra fields')
        if row['period'] != period or row['center'] not in CENTERS or row['code'] not in CODES:
            raise ValueError(f'Unexpected period, center or code: {row}')
        raw = row['total']
        if isinstance(raw, bool) or not re.fullmatch(r'\d+', str(raw)):
            raise ValueError('Totals must be non-negative integers')
        total = int(raw)
        if total > 1_000_000_000:
            raise ValueError('Total exceeds demo limit')
        key = row['center'], row['code']
        if key in seen:
            raise ValueError(f'Duplicate center/code: {key}')
        seen.add(key)
        clean.append((period, *key, total))
    expected = {(center, code) for center in CENTERS for code in CODES}
    if seen != expected:
        raise ValueError(f'Missing center/code pairs: {sorted(expected - seen)}')
    return sorted(clean)


def connect(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(root / 'history.sqlite')
    db.execute('CREATE TABLE IF NOT EXISTS totals (period TEXT, center TEXT, code TEXT, total INTEGER, PRIMARY KEY(period,center,code))')
    db.execute('CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY, at TEXT, action TEXT, period TEXT, reason TEXT, source_sha256 TEXT, before_json TEXT, after_json TEXT, warnings_json TEXT)')
    return db


def repeated_periods(db, rows):
    warnings = []
    period = rows[0][0]
    for center in CENTERS:
        current = [(r[2], r[3]) for r in rows if r[1] == center]
        prior = db.execute('SELECT DISTINCT period FROM totals WHERE center=? AND period<? ORDER BY period', (center, period)).fetchall()
        for (old_period,) in prior:
            previous = db.execute('SELECT code,total FROM totals WHERE center=? AND period=? ORDER BY code', (center, old_period)).fetchall()
            if current == previous:
                warnings.append(f'{center}: {period} matches {old_period} exactly')
    return warnings


def ingest(root, source, period, apply=False, correction=False, reason=''):
    rows = read_input(source, period)
    db = connect(root)
    try:
        db.execute('BEGIN IMMEDIATE')
        before = db.execute('SELECT period,center,code,total FROM totals WHERE period=? ORDER BY center,code', (period,)).fetchall()
        if before == rows:
            return {'status': 'unchanged', 'rows': len(rows), 'total': sum(r[3] for r in rows), 'warnings': []}
        if before and not correction:
            raise ValueError('Period exists; use --correction --reason to replace it')
        if correction and (not before or not reason.strip()):
            raise ValueError('Correction requires an existing period and a non-empty reason')
        warnings = repeated_periods(db, rows)
        result = {'status': 'applied' if apply else 'review', 'rows': len(rows), 'total': sum(r[3] for r in rows), 'warnings': warnings}
        if not apply:
            return result
        # A complete SQLite snapshot is made before any data replacement.
        backups = Path(root) / 'backups'
        backups.mkdir(exist_ok=True)
        backup_path = backups / (datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%f') + '.sqlite')
        with sqlite3.connect(Path(root) / 'history.sqlite') as reader, sqlite3.connect(backup_path) as target:
            reader.backup(target)
        db.execute('DELETE FROM totals WHERE period=?', (period,))
        db.executemany('INSERT INTO totals VALUES (?,?,?,?)', rows)
        db.execute('INSERT INTO audit(at,action,period,reason,source_sha256,before_json,after_json,warnings_json) VALUES(?,?,?,?,?,?,?,?)',
                   (datetime.now(timezone.utc).isoformat(), 'correct' if before else 'load', period, reason.strip(), hashlib.sha256(Path(source).read_bytes()).hexdigest(), json.dumps(before), json.dumps(rows), json.dumps(warnings)))
        db.commit()
        return result
    finally:
        db.close()


def undo(root, audit_id, apply=False):
    db = connect(root)
    try:
        db.execute('BEGIN IMMEDIATE')
        entry = db.execute('SELECT period,before_json,after_json FROM audit WHERE id=?', (audit_id,)).fetchone()
        if not entry:
            raise ValueError('Unknown audit id')
        period, before_json, after_json = entry
        current = db.execute('SELECT period,center,code,total FROM totals WHERE period=? ORDER BY center,code', (period,)).fetchall()
        if current != [tuple(row) for row in json.loads(after_json)]:
            raise ValueError('Period changed since this audit entry; undo refused')
        before = json.loads(before_json)
        if apply:
            db.execute('DELETE FROM totals WHERE period=?', (period,))
            db.executemany('INSERT INTO totals VALUES (?,?,?,?)', before)
            db.execute('INSERT INTO audit(at,action,period,reason,source_sha256,before_json,after_json,warnings_json) VALUES(?,?,?,?,?,?,?,?)',
                       (datetime.now(timezone.utc).isoformat(), 'undo', period, f'Undo audit {audit_id}', '', json.dumps(current), json.dumps(before), '[]'))
            db.commit()
        return {'status': 'undone' if apply else 'review', 'period': period, 'restored_rows': len(before)}
    finally:
        db.close()


def generate(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    rng = random.Random(42026)
    months = {}
    for month in range(1, 4):
        period = f'2026-{month:02}'
        rows = [(period, center, code, rng.randrange(50, 400)) for center in CENTERS for code in CODES]
        if month == 3:
            rows = [(p, c, k, next(r[3] for r in months[1] if r[1:3] == (c,k)) if c == 'Aurora' else n) for p,c,k,n in rows]
        months[month] = rows
        with (directory / f'{period}.csv').open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(FIELDS)
            writer.writerows(rows)
    corrected = [(p,c,k,n+17 if (c,k) == ('Aurora','A01') else n) for p,c,k,n in months[3]]
    with (directory / '2026-03-corrected.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(FIELDS)
        writer.writerows(corrected)


def report(root, year):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    if not 1 <= year <= 9999:
        raise ValueError('Invalid year')
    db = connect(root)
    try:
        rows = db.execute('SELECT period,center,code,total FROM totals WHERE substr(period,1,4)=? ORDER BY period,center,code', (f'{year:04}',)).fetchall()
        audit = db.execute('SELECT id,at,action,period,reason,source_sha256,warnings_json FROM audit ORDER BY id').fetchall()
    finally:
        db.close()
    if not rows:
        raise ValueError('No data for this year')
    out = Path(root) / 'reports'
    out.mkdir(exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Monthly totals'
    sheet.append(['Synthetic monthly operations - demo data only'])
    sheet.append(['Period', 'Aurora', 'Boreal', 'Coral', 'Total'])
    for period in sorted({r[0] for r in rows}):
        sums = [sum(r[3] for r in rows if r[0] == period and r[1] == center) for center in CENTERS]
        sheet.append([period, *sums, sum(sums)])
    sheet.freeze_panes = 'B3'
    sheet.auto_filter.ref = f'A2:E{sheet.max_row}'
    for cell in sheet[2]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='153D52')
    for col in 'ABCDE':
        sheet.column_dimensions[col].width = 22
    chart = BarChart()
    chart.title = 'Monthly volume by fictional center'
    chart.add_data(Reference(sheet, min_col=2, max_col=4, min_row=2, max_row=sheet.max_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=sheet.max_row))
    sheet.add_chart(chart, 'G2')
    raw = book.create_sheet('Records')
    raw.append(FIELDS)
    for row in rows:
        raw.append(row)
    log = book.create_sheet('Audit')
    log.append(['ID', 'UTC timestamp', 'Action', 'Period', 'Reason', 'Source SHA256', 'Warnings'])
    for row in audit:
        # Escape spreadsheet formula prefixes in free text.
        log.append(["'"+v if isinstance(v,str) and v.startswith(('=', '+', '-', '@')) else v for v in row])
    book.save(out / f'annual-{year}.xlsx')
    for period in sorted({r[0] for r in rows}):
        for center in CENTERS:
            pdf = canvas.Canvas(str(out / f'{period}-{center}.pdf'), pagesize=A4)
            pdf.setTitle(f'Synthetic receipt {period} {center}')
            pdf.setFont('Helvetica-Bold', 20)
            pdf.drawString(50, 785, 'Monthly Operations')
            pdf.setFont('Helvetica', 11)
            pdf.drawString(50, 758, 'SYNTHETIC DATA / PORTFOLIO DEMO')
            pdf.drawString(50, 725, f'Period: {period}    Fictional center: {center}')
            y = 680
            selected = [r for r in rows if r[0] == period and r[1] == center]
            for _,_,code,total in selected:
                pdf.drawString(65,y,f'{code}: {total:,}')
                y -= 28
            pdf.setFont('Helvetica-Bold', 13)
            pdf.drawString(65,y-10,f'Total: {sum(r[3] for r in selected):,}')
            pdf.setFont('Helvetica', 9)
            pdf.drawString(50, 65, 'Fictional receipt. Not an official record. Regenerate after corrections or undo.')
            pdf.save()
    (out / 'audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
    return str(out)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workdir', default='work')
    sub = parser.add_subparsers(dest='command', required=True)
    gen = sub.add_parser('generate')
    gen.add_argument('--output', default='examples')
    load = sub.add_parser('load')
    load.add_argument('source')
    load.add_argument('--period', required=True)
    load.add_argument('--apply', action='store_true')
    load.add_argument('--correction', action='store_true')
    load.add_argument('--reason', default='')
    rev = sub.add_parser('undo')
    rev.add_argument('audit_id', type=int)
    rev.add_argument('--apply', action='store_true')
    rep = sub.add_parser('report')
    rep.add_argument('--year', type=int, required=True)
    args = parser.parse_args()
    try:
        if args.command == 'generate':
            generate(args.output)
            result = {'generated': args.output, 'seed': 42026}
        elif args.command == 'load':
            result = ingest(args.workdir,args.source,args.period,args.apply,args.correction,args.reason)
        elif args.command == 'undo':
            result = undo(args.workdir,args.audit_id,args.apply)
        else:
            result = {'reports': report(args.workdir,args.year)}
        print(json.dumps(result, indent=2))
    except (ValueError, OSError, sqlite3.Error) as error:
        parser.exit(2, f'Error: {error}\n')


if __name__ == '__main__':
    main()
