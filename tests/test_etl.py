import csv
import sqlite3
import tempfile
import unittest
from pathlib import Path

from etl import generate, ingest, read_input, report, undo


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'inputs'
        self.work = self.root / 'work'
        generate(self.inputs)

    def load(self, month, **kwargs):
        return ingest(self.work, self.inputs / f'2026-{month:02}.csv', f'2026-{month:02}', **kwargs)

    def records(self):
        with sqlite3.connect(self.work / 'history.sqlite') as db:
            return db.execute('SELECT * FROM totals ORDER BY period,center,code').fetchall()

    def test_review_does_not_load(self):
        self.assertEqual(self.load(1)['status'], 'review')
        self.assertEqual(self.records(), [])

    def test_idempotency(self):
        self.load(1, apply=True)
        self.assertEqual(self.load(1, apply=True)['status'], 'unchanged')
        with sqlite3.connect(self.work / 'history.sqlite') as db:
            self.assertEqual(db.execute('SELECT count(*) FROM audit').fetchone()[0], 1)

    def test_nonadjacent_duplicate_is_warning(self):
        for month in (1,2):
            self.load(month, apply=True)
        result = self.load(3, apply=True)
        self.assertEqual(result['warnings'], ['Aurora: 2026-03 matches 2026-01 exactly'])
        self.assertEqual(len(self.records()), 36)

    def test_correction_preserves_other_months_and_undo(self):
        for month in (1,2,3):
            self.load(month, apply=True)
        before = self.records()
        corrected = self.inputs / '2026-03-corrected.csv'
        with self.assertRaises(ValueError):
            ingest(self.work, corrected, '2026-03', True)
        with self.assertRaises(ValueError):
            ingest(self.work, corrected, '2026-03', True, True)
        ingest(self.work, corrected, '2026-03', True, True, 'Synthetic correction')
        after = self.records()
        self.assertEqual(before[:24], after[:24])
        self.assertEqual(sum(r[3] for r in after)-sum(r[3] for r in before),17)
        undo(self.work,4,True)
        self.assertEqual(before,self.records())
        self.assertEqual(len(list((self.work/'backups').glob('*.sqlite'))),4)

    def test_stale_undo_refused(self):
        self.load(3, apply=True)
        ingest(self.work,self.inputs/'2026-03-corrected.csv','2026-03',True,True,'Revision')
        with self.assertRaises(ValueError):
            undo(self.work,1,True)

    def test_reject_missing_duplicate_negative_and_bad_period(self):
        source = self.inputs/'2026-01.csv'
        with source.open(newline='') as stream:
            data = list(csv.reader(stream))
        cases = [data[:-1], data+[data[1]], [row[:] for row in data], [row[:] for row in data]]
        cases[2][1][3] = '-1'
        cases[3][1][0] = '2026-13'
        for index, rows in enumerate(cases):
            target = self.inputs/f'invalid{index}.csv'
            with target.open('w',newline='') as stream:
                csv.writer(stream).writerows(rows)
            with self.assertRaises(ValueError):
                read_input(target,'2026-01')

    def test_xlsx_input(self):
        from openpyxl import Workbook
        source = self.inputs/'2026-01.csv'
        book = Workbook()
        with source.open(newline='') as stream:
            for index,row in enumerate(csv.reader(stream)):
                book.active.append(row if index == 0 else row[:3]+[int(row[3])])
        target = self.inputs/'input.xlsx'
        book.save(target)
        self.assertEqual(read_input(source,'2026-01'),read_input(target,'2026-01'))

    def test_reports_reconcile(self):
        from openpyxl import load_workbook
        for month in (1,2,3):
            self.load(month,apply=True)
        out = Path(report(self.work,2026))
        book = load_workbook(out/'annual-2026.xlsx',data_only=True)
        try:
            summary_total = sum(row[4] for row in book['Monthly totals'].iter_rows(min_row=3,values_only=True))
            self.assertEqual(summary_total,sum(r[3] for r in self.records()))
            self.assertEqual(book['Records'].max_row,37)
        finally:
            book.close()
        self.assertEqual(len(list(out.glob('*.pdf'))),9)
        self.assertTrue((out/'audit.json').exists())

    def test_generator_reproducible(self):
        other = self.root/'other'
        generate(other)
        for source in self.inputs.glob('*.csv'):
            self.assertEqual(source.read_bytes(),(other/source.name).read_bytes())


if __name__ == '__main__':
    unittest.main()
